"""导出 Excel：两种账单各一套写表函数，单元格样式共用。"""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from detail import Detail, extra_columns, summaries_of
from statement import Group, Page, Record, all_records

HEAD = ["序号", "日期", "交易类型", "金额(元)", "卡/渠道", "时间", "交易后余额(元)", "来源工作表"]
WIDTHS = [6, 12, 11, 13, 16, 8, 16, 16]

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_TOTAL_FONT = Font(bold=True, color="1F4E79", size=10)
_OK_FILL = PatternFill("solid", fgColor="E2EFDA")
_OK_FONT = Font(bold=True, color="375623")
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
_BAD_FONT = Font(bold=True, color="9C0006")
_HIT_FILL = PatternFill("solid", fgColor="FFD400")
_HIT_FONT = Font(bold=True, color="7F3F00")
_MARK_FILL = PatternFill("solid", fgColor="FF9999")
"""用户在截图上用红框圈出来的行 —— 需要特别核对，红色底纹。"""
_MARK_FONT = Font(bold=True, color="7F0000")
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")
_MONEY = "#,##0.00"


def _headers(ws, row: int, columns: Sequence[str]) -> None:
    for index, name in enumerate(columns, 1):
        cell = ws.cell(row=row, column=index, value=name)
        cell.fill, cell.font, cell.alignment, cell.border = _HEAD_FILL, _HEAD_FONT, _CENTER, _BORDER


def _write_records(ws, records: Sequence[Record], start: int, with_source: bool) -> int:
    """写入明细表，返回最后一行的行号。"""
    _headers(ws, start, HEAD[: len(HEAD) if with_source else len(HEAD) - 1])
    row = start
    for index, record in enumerate(records, 1):
        row += 1
        values = [
            index,
            record.date or "未知",
            record.kind,
            record.amount,
            record.card,
            record.time,
            record.balance,
            record.source,
        ]
        if not with_source:
            values.pop()
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            if column in (1, 2, 3, 5, 6, 8):
                cell.alignment = _CENTER
            else:
                cell.alignment = _RIGHT
                if column in (4, 7) and isinstance(value, (int, float)):
                    cell.number_format = _MONEY
    return row


def _highlight(ws, records: Sequence[Record], start: int, hits: set[float], with_source: bool) -> int:
    """给命中金额的整行加黄底，返回高亮行数。"""
    count = 0
    last_column = len(HEAD) if with_source else len(HEAD) - 1
    for offset, record in enumerate(records, 1):
        if round(abs(record.amount), 2) not in hits:
            continue
        count += 1
        row = start + offset
        for column in range(1, last_column + 1):
            ws.cell(row=row, column=column).fill = _HIT_FILL
        for column in (3, 4):
            ws.cell(row=row, column=column).font = _HIT_FONT
    return count


def _mark_flagged(ws, records: Sequence[Record], start: int, with_source: bool) -> int:
    """给用户红框圈住的行加红底，返回标记行数。

    要在 ``_highlight`` **之后**调用 —— 红框是用户亲手标的，比「金额命中」更具体，
    两种都命中时红色应当覆盖黄色。
    """
    count = 0
    last_column = len(HEAD) if with_source else len(HEAD) - 1
    for offset, record in enumerate(records, 1):
        if not record.flagged:
            continue
        count += 1
        row = start + offset
        for column in range(1, last_column + 1):
            ws.cell(row=row, column=column).fill = _MARK_FILL
        for column in (3, 4):
            ws.cell(row=row, column=column).font = _MARK_FONT
    return count


def _totals(records: Sequence[Record]) -> tuple[float, float]:
    expense = round(sum(-r.amount for r in records if r.amount < 0), 2)
    income = round(sum(r.amount for r in records if r.amount > 0), 2)
    return expense, income


def _unique_sheet_name(name: str, used: set[str]) -> str:
    """同名工作表加序号（Excel 不允许重名，且名字上限 31 字符）。"""
    candidate = name[:31] or "明细"
    index = 2
    while candidate in used:
        suffix = f"({index})"
        candidate = f"{name[: 31 - len(suffix)]}{suffix}"
        index += 1
    return candidate


def write_statement(
    out_path: Path,
    groups: Sequence[Group],
    highlight: Iterable[float],
    pages: Sequence[Page] | None = None,
) -> Path:
    """生成「好分期账单长截图」的 Excel。

    Args:
        out_path: 输出文件路径。
        groups: 按账单区间归组并去重后的结果。
        highlight: 需要高亮的金额（按绝对值匹配）。
        pages: 全部页；为 None 时从 groups 里取。

    Returns:
        实际写入的文件路径。

    Raises:
        PermissionError: 目标文件被 Excel 等程序占用。
    """
    hits = {round(float(v), 2) for v in highlight}
    rank = {page.source: (order, position) for order, group in enumerate(groups)
            for position, page in enumerate(group.pages)}
    candidates: list[Page] = list(pages) if pages is not None else [p for g in groups for p in g.pages]
    # 工作表顺序跟着账单时间走，不跟着文件名
    all_pages = sorted(candidates, key=lambda p: rank.get(p.source, (len(groups), 0)))
    records = all_records(groups)
    total_expense, total_income = _totals(records)
    # 红框标记的数量要在写「核对说明」之前就知道，所以先数一遍
    marked_count = sum(1 for record in records if record.flagged)
    file_count = len(all_pages) + sum(len(p.copies) for p in all_pages)

    workbook = Workbook()

    # ---------- 汇总 ----------
    ws = workbook.active
    ws.title = "汇总"
    ws["A1"] = (
        f"账单明细汇总：{file_count} 个图片文件 → 按内容去重后 {len(all_pages)} 张唯一截图 "
        f"→ {len(groups)} 个账单期间，共 {len(records)} 笔交易"
    )
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")

    summary_head = [
        "年份",
        "账单区间",
        "数据来源",
        "交易笔数",
        "支出合计(元)",
        "收入合计(元)",
        "净收入(元)",
        "页面显示支出(元)",
        "支出差异(元)",
        "核对结果",
    ]
    _headers(ws, 3, summary_head)

    row = 3
    for group in groups:
        row += 1
        diff = group.expense_diff
        if diff is None:
            verdict, verdict_fill, verdict_font = "无汇总卡", _TOTAL_FILL, _TOTAL_FONT
        elif abs(diff) < 0.005:
            verdict, verdict_fill, verdict_font = "一致", _OK_FILL, _OK_FONT
        else:
            verdict, verdict_fill, verdict_font = f"差 {diff:,.2f}", _BAD_FILL, _BAD_FONT
        sources = "、".join(p.source for p in group.pages)
        values = [
            group.year_label,
            group.period_label,
            sources,
            len(group.records),
            group.expense,
            group.income,
            round(group.income - group.expense, 2),
            group.card_expense,
            diff,
            verdict,
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if column == 3 else (_CENTER if column in (1, 4, 10) else _RIGHT)
            if column in (5, 6, 7, 8, 9) and isinstance(value, (int, float)):
                cell.number_format = _MONEY
        for column in (9, 10):
            ws.cell(row=row, column=column).fill = verdict_fill
            ws.cell(row=row, column=column).font = verdict_font

    row += 1
    card_totals = [g.card_expense for g in groups if g.card_expense is not None]
    card_expense = round(sum(card_totals), 2) if card_totals else None
    total_diff = round(total_expense - card_expense, 2) if card_expense is not None else None
    if total_diff is None:
        total_verdict = "无汇总卡"
    elif abs(total_diff) < 0.005:
        total_verdict = "一致"
    else:
        total_verdict = f"差 {total_diff:,.2f}"
    values = [
        "合计",
        "",
        f"{len(groups)} 个期间 / {len(all_pages)} 张截图",
        len(records),
        total_expense,
        total_income,
        round(total_income - total_expense, 2),
        card_expense,
        total_diff,
        total_verdict,
    ]
    for column, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.border, cell.fill, cell.font = _BORDER, _TOTAL_FILL, _TOTAL_FONT
        cell.alignment = _LEFT if column == 3 else (_CENTER if column in (1, 4, 10) else _RIGHT)
        if column in (5, 6, 7, 8, 9) and isinstance(value, (int, float)):
            cell.number_format = _MONEY

    row += 2
    notes = [
        "核对说明",
        "1. 页面顶部汇总卡的「支出/收入」是该搜索条件下的全额合计，这里与明细逐笔求和比对；差异为 0.00 即完全一致。",
        "2. 差异不为 0 时已标红，逐组原因见下方「缺口说明」。",
        f"3. 黄色底纹 = 金额为 {_format_hits(hits)} 元的交易。",
        f"4. 红色底纹 = 你在截图上用红框圈出来、需要特别核对的行（共 {marked_count} 笔）。",
        "5. 跨图去重：同一笔交易可能同时出现在长截图和单屏截图里，已按「日期+类型+金额+余额」合并，见下方「跨图重复」。",
        "6. 图片文件按内容 MD5 去重，内容完全相同的副本只算一张，见下方文件对照。",
        "7. 定位方式是自校准：从每张图自己的金额块量出行距与字号，不写死分辨率，换手机/换截图工具都不受影响。",
    ]
    for note in notes:
        ws.cell(row=row, column=1, value=note).font = Font(size=9, bold=note == "核对说明")
        row += 1

    gaps = [(g.year_label, g.gap_note) for g in groups if g.gap_note]
    if gaps:
        row += 1
        ws.cell(row=row, column=1, value="缺口说明").font = Font(bold=True, size=10, color="1F4E79")
        row += 1
        for label, note in gaps:
            ws.cell(row=row, column=1, value=f"{label}：{note}").font = Font(size=9)
            row += 1

    crosses = [(kept, dropped, count) for g in groups for kept, dropped, count in g.cross]
    if crosses:
        row += 1
        ws.cell(row=row, column=1, value="跨图重复（已合并）").font = Font(
            bold=True, size=10, color="1F4E79"
        )
        row += 1
        for kept, dropped, count in crosses:
            ws.cell(row=row, column=1,
                    value=f"{dropped} 有 {count} 笔与 {kept} 重复，已只保留一笔。").font = Font(size=9)
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="文件对照（重复文件 → 唯一截图）").font = Font(
        bold=True, size=10, color="1F4E79"
    )
    row += 1
    _headers(ws, row, ["工作表", "账单区间", "代表文件", "重复副本"])
    for page in all_pages:
        row += 1
        group = next((g for g in groups if page in g.pages), None)
        values = [
            page.name,
            group.period_label if group else "",
            page.source,
            "、".join(page.copies) or "—",
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = _CENTER if column == 1 else _LEFT

    for index, width in enumerate([12, 26, 34, 14, 15, 15, 14, 18, 13, 14], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"

    # ---------- 全部明细 ----------
    ws_all = workbook.create_sheet("全部明细(去重)")
    ws_all["A1"] = (
        f"全部交易明细（去重后 {len(records)} 笔；支出 {total_expense:,.2f} 元，收入 {total_income:,.2f} 元）"
    )
    ws_all["A1"].font = Font(bold=True, size=11, color="1F4E79")
    end = _write_records(ws_all, records, 2, with_source=True)
    _highlight(ws_all, records, 2, hits, with_source=True)
    _mark_flagged(ws_all, records, 2, with_source=True)
    row = end + 2
    ws_all.cell(row=row, column=2, value="支出合计").font = Font(bold=True)
    cell = ws_all.cell(row=row, column=4, value=total_expense)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    ws_all.cell(row=row, column=5, value="收入合计").font = Font(bold=True)
    cell = ws_all.cell(row=row, column=6, value=total_income)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    for index, width in enumerate(WIDTHS, 1):
        ws_all.column_dimensions[get_column_letter(index)].width = width
    ws_all.freeze_panes = "A3"

    # ---------- 每张唯一截图一个表 ----------
    used_names = {ws.title, ws_all.title}
    for page in all_pages:
        sheet_name = _unique_sheet_name(page.name, used_names)
        used_names.add(sheet_name)
        ws_page = workbook.create_sheet(sheet_name)
        page_records = sorted(page.records, key=lambda r: (r.date, r.time))
        expense, income = _totals(page_records)
        group = next((g for g in groups if page in g.pages), None)
        ws_page["A1"] = (
            f"{page.name}　来源：{page.source}"
            + (f"（副本 {'、'.join(page.copies)}）" if page.copies else "")
            + (f"　区间：{group.period_label}" if group else "")
        )
        ws_page["A1"].font = Font(bold=True, size=11, color="1F4E79")
        ws_page["A2"] = (
            f"共 {len(page_records)} 笔　支出 {expense:,.2f} 元　收入 {income:,.2f} 元"
            f"　黄色行为 {_format_hits(hits)} 元"
            + (f"　本图有 {page.marks} 个红框标注（红底行）" if page.marks else "")
        )
        ws_page["A2"].font = Font(size=9, color="7F3F00", bold=True)
        _write_records(ws_page, page_records, 3, with_source=False)
        _highlight(ws_page, page_records, 3, hits, with_source=False)
        _mark_flagged(ws_page, page_records, 3, with_source=False)
        for index, width in enumerate(WIDTHS[:-1], 1):
            ws_page.column_dimensions[get_column_letter(index)].width = width
        ws_page.freeze_panes = "A4"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(out_path)
    except PermissionError as exc:  # 文件被 Excel 打开时 openpyxl 抛的就是这个
        raise PermissionError(f"无法写入 {out_path.name}，请先关闭同名文件") from exc
    return out_path


def _format_hits(hits: Iterable[float]) -> str:
    values = sorted({float(v) for v in hits})
    return " / ".join(f"{v:g}" for v in values) if values else "（无）"


# ========================= 收支详情单页 =========================

DETAIL_HEAD = [
    "序号", "交易日期", "交易时间", "金额(元)", "业务摘要", "交易场所",
    "交易后余额(元)", "交易卡号", "交易账户", "交易户名", "源文件",
]
DETAIL_WIDTHS = [6, 12, 10, 12, 10, 30, 16, 15, 15, 10, 22]

EXTRA_WIDTHS = {"交易流水号": 36, "对方账户": 30, "商品信息": 26, "商户名称": 22}


def _extra_width(label: str) -> float:
    """版式特有字段的列宽：流水号这种长值要给够。"""
    return EXTRA_WIDTHS.get(label, 16)



def _section(ws, row: int, title: str, columns: Sequence[str], rows: Sequence[Sequence],
             money_cols: Sequence[int] = (), percent_cols: Sequence[int] = ()) -> int:
    """写一个「小标题 + 表头 + 数据」区块，返回下一个可用行号。"""
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=11, color="1F4E79")
    row += 1
    _headers(ws, row, columns)
    row += 1
    for values in rows:
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = _RIGHT if column in money_cols or column in percent_cols else _CENTER
            if column in money_cols and isinstance(value, (int, float)):
                cell.number_format = _MONEY
            if column in percent_cols and isinstance(value, (int, float)):
                cell.number_format = "0.0%"
        row += 1
    return row + 1


def _group_rows(keys: Sequence[str], amounts: Sequence[float], total: float) -> list[list]:
    """按 key 分组统计笔数/金额/占比。"""
    counts: Counter[str] = Counter(keys)
    sums: dict[str, float] = {}
    for key, value in zip(keys, amounts):
        sums[key] = round(sums.get(key, 0.0) + abs(value), 2)
    return [
        [key, counts[key], sums[key], (sums[key] / total if total else 0.0)]
        for key in sorted(counts)
    ]


def write_detail(
    out_path: Path,
    details: Sequence[Detail],
    highlight: Iterable[float],
    file_count: int = 0,
    duplicates: dict[str, str] | None = None,
    repeats: dict[str, str] | None = None,
) -> Path:
    """生成「收支详情单页」的 Excel：明细 + 多维汇总。

    Args:
        out_path: 输出文件路径。
        details: 去重后的一笔笔交易。
        highlight: 需要高亮的金额（按绝对值匹配）。
        file_count: 目录里的图片文件总数（用于说明去重情况）。
        duplicates: 内容完全相同的副本：副本文件名 -> 代表文件名。
        repeats: 同一笔交易被多张不同截图记录：重复截图名 -> 代表文件名。

    Returns:
        实际写入的文件路径。

    Raises:
        PermissionError: 目标文件被 Excel 等程序占用。
        ValueError: 一笔交易都没有。
    """
    if not details:
        raise ValueError("没有解析出任何交易")
    hits = {round(float(v), 2) for v in highlight}
    report = summaries_of(details)
    ordered = sorted(details, key=lambda d: d.time or "9999")
    amounts = [d.signed_amount or 0.0 for d in ordered]
    expense = round(sum(-v for v in amounts if v < 0), 2)
    income = round(sum(v for v in amounts if v > 0), 2)
    total = round(expense + income, 2)
    duplicates = duplicates or {}
    repeats = repeats or {}
    file_count = file_count or len(ordered) + len(duplicates) + len(repeats)
    unique_images = file_count - len(duplicates)

    workbook = Workbook()

    # ---------- 明细 ----------
    ws = workbook.active
    ws.title = "明细"
    ws["A1"] = (
        f"交易明细：{file_count} 个图片文件 → 内容去重 {unique_images} 张"
        f" → 同一笔合并后 {len(ordered)} 笔交易"
        + (f"　版式：{'、'.join(report['profiles'])}" if report["profiles"] else "")
    )
    ws["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws["A2"] = (
        f"支出 {expense:,.2f} 元　收入 {income:,.2f} 元　"
        f"黄色行为 {_format_hits(hits)} 元"
    )
    ws["A2"].font = Font(size=9, color="7F3F00", bold=True)

    # 版式特有字段按实际出现过的标签动态加列（两种版式的列不完全一样）
    extras = extra_columns(ordered)
    columns = list(DETAIL_HEAD) + extras
    _headers(ws, 3, columns)
    row = 3
    for index, item in enumerate(ordered, 1):
        row += 1
        values = [
            index, item.date, item.clock, item.signed_amount, item.summary, item.place,
            item.balance, item.card, item.account, item.holder, item.source,
        ] + [item.extra.get(label, "") for label in extras]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            if column in (4, 7):
                cell.alignment = _RIGHT
                if isinstance(value, (int, float)):
                    cell.number_format = _MONEY
            elif column > len(DETAIL_HEAD):
                cell.alignment = _LEFT
            else:
                cell.alignment = _CENTER
        if item.signed_amount is not None and round(abs(item.signed_amount), 2) in hits:
            for column in range(1, len(values) + 1):
                ws.cell(row=row, column=column).fill = _HIT_FILL
            ws.cell(row=row, column=4).font = _HIT_FONT
    row += 1
    ws.cell(row=row, column=2, value="支出合计").font = Font(bold=True)
    cell = ws.cell(row=row, column=4, value=expense)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    ws.cell(row=row, column=5, value="收入合计").font = Font(bold=True)
    cell = ws.cell(row=row, column=6, value=income)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    for index, width in enumerate(DETAIL_WIDTHS + [_extra_width(l) for l in extras], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"

    # ---------- 汇总 ----------
    ws2 = workbook.create_sheet("汇总")
    ws2["A1"] = f"收支详情多维汇总：{len(ordered)} 笔交易"
    ws2["A1"].font = Font(bold=True, size=13, color="1F4E79")
    row = 3

    row = _section(ws2, row, "总览", ["项目", "数值"], [
        ["交易笔数", len(ordered)],
        ["支出合计(元)", expense],
        ["收入合计(元)", income],
        ["净额(元)", round(income - expense, 2)],
        ["时间范围", f"{ordered[0].date} ~ {ordered[-1].date}"],
        ["识别到的版式", "、".join(report["profiles"]) or "—"],
        ["交易卡号", "、".join(report["cards"]) or "—"],
        ["交易账户", "、".join(report["accounts"]) or "—"],
        ["交易户名", "、".join(report["holders"]) or "—"],
    ], money_cols=(2,))

    row = _section(ws2, row, "按年",
                   ["年份", "笔数", "金额(元)", "占比"],
                   _group_rows([d.year for d in ordered], amounts, total),
                   money_cols=(3,), percent_cols=(4,))
    row = _section(ws2, row, "按月",
                   ["年月", "笔数", "金额(元)", "占比"],
                   _group_rows([d.month for d in ordered], amounts, total),
                   money_cols=(3,), percent_cols=(4,))
    row = _section(
        ws2, row, "按金额档", ["金额(元)", "笔数", "小计(元)", "占比"],
        _group_rows([round(d.signed_amount or 0.0, 2) for d in ordered], amounts, total),
        money_cols=(1, 3), percent_cols=(4,),
    )
    row = _section(ws2, row, "按业务摘要", ["业务摘要", "笔数", "金额(元)", "占比"],
                   _group_rows([d.summary or "未识别" for d in ordered], amounts, total),
                   money_cols=(3,), percent_cols=(4,))
    row = _section(ws2, row, "按交易场所", ["交易场所", "笔数", "金额(元)", "占比"],
                   _group_rows([d.place or "未识别" for d in ordered], amounts, total),
                   money_cols=(3,), percent_cols=(4,))

    checked = len(ordered) - len(report["unknown"])
    if checked == 0:
        amount_check = "该版式没有「交易金额」字段，跳过（改看下一行的日期交叉校验）"
    else:
        amount_check = (
            f"{checked}/{len(ordered)} 通过；不一致 {len(report['inconsistent'])} 张"
        )
    date_checked = report["date_checked"]
    if date_checked == 0:
        date_check = "该版式没有「记账日」字段，跳过"
    else:
        date_check = (
            f"{date_checked}/{len(ordered)} 通过；不一致 {len(report['date_inconsistent'])} 张"
        )
    issues: list[list] = [
        ["交易金额 与 顶部金额 一致", amount_check],
        ["记账日 与 交易时间日期 一致", date_check],
        ["未读到的字段", "无" if not report["missing"] else f"{len(report['missing'])} 张（明细见下）"],
        ["重复图片文件（内容相同）", f"{len(duplicates)} 个（已去重）"],
        ["重复截图（同一笔交易）", f"{len(repeats)} 张（已合并）"],
        ["原始文件 → 去重后",
         f"{file_count} 个 → {unique_images} 张唯一截图 → {len(ordered)} 笔交易"],
    ]
    for source in report["inconsistent"]:
        issues.append(["金额不一致截图", source])
    for source in report["date_inconsistent"]:
        issues.append(["日期不一致截图", source])
    for source, fields in report["missing"].items():
        issues.append(["未读到的字段", f"{source}（缺 {'、'.join(fields)}）"])
    for copy_name, source in duplicates.items():
        issues.append(["重复文件（内容相同）", f"{copy_name} = {source}"])
    for copy_name, source in repeats.items():
        issues.append(["重复截图（同一笔）", f"{copy_name} = {source}"])
    row = _section(ws2, row, "自动校验", ["项目", "结果"], issues)

    ws2.cell(row=row, column=1, value=(
        "说明：① 版式按图上命中的特征标签自动识别，两种版式（收支详情单页 / 建行「明细详情」）"
        "共用同一套「标签—值」取值逻辑；② 建行版式没有「交易金额」字段，"
        "改用「记账日」与「交易时间」的日期是否一致来交叉校验；"
        "③ 内容完全相同的图片按 MD5 去重；④ 交易时间 + 金额 + 余额 + 账户都相同视为同一笔交易"
        "被重复截图，已合并；⑤ 「交易流水号」「对方账户」这类换行的长值会按行拼接后写入；"
        "⑥ 报「未读到的字段」不一定是识别问题 —— 部分交易在页面上本来就没有某些行"
        "（例如建行版式的「商品信息」），属正常差异。"
    )).font = Font(size=9)
    for index, width in enumerate([34, 22, 16, 12], 1):
        ws2.column_dimensions[get_column_letter(index)].width = width

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(out_path)
    except PermissionError as exc:  # 文件被 Excel 打开时 openpyxl 抛的就是这个
        raise PermissionError(f"无法写入 {out_path.name}，请先关闭同名文件") from exc
    return out_path
