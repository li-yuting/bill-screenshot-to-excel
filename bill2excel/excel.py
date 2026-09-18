"""导出 Excel：汇总表 + 全部明细 + 每张唯一截图一个工作表。"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from parse import Group, Page, Record, all_records

HEAD = ["序号", "日期", "交易类型", "金额(元)", "卡/渠道", "时间", "交易后余额(元)", "来源工作表"]
WIDTHS = [6, 12, 11, 13, 16, 8, 16, 16]

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_TOTAL_FONT = Font(bold=True, color="1F4E79", size=10)
_OK_FILL = PatternFill("solid", fgColor="E2EFDA")
_OK_FONT = Font(bold=True, color="375623")
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
_BAD_FONT = Font(bold=True, color="9C0006")
_HIT_FILL = PatternFill("solid", fgColor="FFD400")
_HIT_FONT = Font(bold=True, color="7F3F00")
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")
_MONEY = "#,##0.00"


def _headers(ws, row: int, columns: Sequence[str]) -> None:
    for index, name in enumerate(columns, 1):
        cell = ws.cell(row=row, column=index, value=name)
        cell.fill, cell.font, cell.alignment, cell.border = _HEAD_FILL, _HEAD_FONT, _CENTER, _BORDER


def _write_records(ws, records: Sequence[Record], start: int, with_source: bool) -> int:
    """写入明细表，返回最后一行的行号。"""
    _headers(ws, start, HEAD[: len(HEAD) if with_source else len(HEAD) - 1])
    row = start
    for index, record in enumerate(records, 1):
        row += 1
        values = [
            index,
            record.date or "未知",
            record.kind,
            record.amount,
            record.card,
            record.time,
            record.balance,
            record.source,
        ]
        if not with_source:
            values.pop()
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            if column in (1, 2, 3, 5, 6, 8):
                cell.alignment = _CENTER
            else:
                cell.alignment = _RIGHT
                if column in (4, 7) and isinstance(value, (int, float)):
                    cell.number_format = _MONEY
    return row


def _highlight(ws, records: Sequence[Record], start: int, hits: set[float], with_source: bool) -> int:
    """给命中金额的整行加黄底，返回高亮行数。"""
    count = 0
    last_column = len(HEAD) if with_source else len(HEAD) - 1
    for offset, record in enumerate(records, 1):
        if round(abs(record.amount), 2) not in hits:
            continue
        count += 1
        row = start + offset
        for column in range(1, last_column + 1):
            ws.cell(row=row, column=column).fill = _HIT_FILL
        for column in (3, 4):
            ws.cell(row=row, column=column).font = _HIT_FONT
    return count


def _totals(records: Sequence[Record]) -> tuple[float, float]:
    expense = round(sum(-r.amount for r in records if r.amount < 0), 2)
    income = round(sum(r.amount for r in records if r.amount > 0), 2)
    return expense, income


def _unique_sheet_name(name: str, used: set[str]) -> str:
    """同名工作表加序号（Excel 不允许重名，且名字上限 31 字符）。"""
    candidate = name[:31] or "明细"
    index = 2
    while candidate in used:
        suffix = f"({index})"
        candidate = f"{name[: 31 - len(suffix)]}{suffix}"
        index += 1
    return candidate


def write_workbook(
    out_path: Path,
    groups: Sequence[Group],
    highlight: Iterable[float],
    pages: Sequence[Page] | None = None,
) -> Path:
    """生成账单 Excel。

    Args:
        out_path: 输出文件路径。
        groups: 按账单区间归组并去重后的结果。
        highlight: 需要高亮的金额（按绝对值匹配）。
        pages: 全部页；为 None 时从 groups 里取。

    Returns:
        实际写入的文件路径。

    Raises:
        PermissionError: 目标文件被 Excel 等程序占用。
    """
    hits = {round(float(v), 2) for v in highlight}
    rank = {page.source: (order, position) for order, group in enumerate(groups)
            for position, page in enumerate(group.pages)}
    candidates: list[Page] = list(pages) if pages is not None else [p for g in groups for p in g.pages]
    # 工作表顺序跟着账单时间走，不跟着文件名
    all_pages = sorted(candidates, key=lambda p: rank.get(p.source, (len(groups), 0)))
    records = all_records(groups)
    total_expense, total_income = _totals(records)
    file_count = len(all_pages) + sum(len(p.copies) for p in all_pages)

    workbook = Workbook()

    # ---------- 汇总 ----------
    ws = workbook.active
    ws.title = "汇总"
    ws["A1"] = (
        f"账单明细汇总：{file_count} 个图片文件 → 按内容去重后 {len(all_pages)} 张唯一截图 "
        f"→ {len(groups)} 个账单期间，共 {len(records)} 笔交易"
    )
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")

    summary_head = [
        "年份",
        "账单区间",
        "数据来源",
        "交易笔数",
        "支出合计(元)",
        "收入合计(元)",
        "净收入(元)",
        "页面显示支出(元)",
        "支出差异(元)",
        "核对结果",
    ]
    _headers(ws, 3, summary_head)

    row = 3
    for group in groups:
        row += 1
        diff = group.expense_diff
        if diff is None:
            verdict, verdict_fill, verdict_font = "无汇总卡", _TOTAL_FILL, _TOTAL_FONT
        elif abs(diff) < 0.005:
            verdict, verdict_fill, verdict_font = "一致", _OK_FILL, _OK_FONT
        else:
            verdict, verdict_fill, verdict_font = f"差 {diff:,.2f}", _BAD_FILL, _BAD_FONT
        sources = "、".join(p.source for p in group.pages)
        values = [
            f"{group.year}年",
            group.period_label,
            sources,
            len(group.records),
            group.expense,
            group.income,
            round(group.income - group.expense, 2),
            group.card_expense,
            diff,
            verdict,
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if column == 3 else (_CENTER if column in (1, 4, 10) else _RIGHT)
            if column in (5, 6, 7, 8, 9) and isinstance(value, (int, float)):
                cell.number_format = _MONEY
        for column in (9, 10):
            ws.cell(row=row, column=column).fill = verdict_fill
            ws.cell(row=row, column=column).font = verdict_font

    row += 1
    card_totals = [g.card_expense for g in groups if g.card_expense is not None]
    card_expense = round(sum(card_totals), 2) if card_totals else None
    total_diff = round(total_expense - card_expense, 2) if card_expense is not None else None
    if total_diff is None:
        total_verdict = "无汇总卡"
    elif abs(total_diff) < 0.005:
        total_verdict = "一致"
    else:
        total_verdict = f"差 {total_diff:,.2f}"
    values = [
        "合计",
        "",
        f"{len(groups)} 个期间 / {len(all_pages)} 张截图",
        len(records),
        total_expense,
        total_income,
        round(total_income - total_expense, 2),
        card_expense,
        total_diff,
        total_verdict,
    ]
    for column, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.border, cell.fill, cell.font = _BORDER, _TOTAL_FILL, _TOTAL_FONT
        cell.alignment = _LEFT if column == 3 else (_CENTER if column in (1, 4, 10) else _RIGHT)
        if column in (5, 6, 7, 8, 9) and isinstance(value, (int, float)):
            cell.number_format = _MONEY

    row += 2
    notes = [
        "核对说明",
        "1. 页面顶部汇总卡的「支出/收入」是该搜索条件下的全额合计，这里与明细逐笔求和比对；差异为 0.00 即完全一致。",
        "2. 差异不为 0 时已标红：长截图拼接可能漏抓个别交易，建议补截该期间账单核对（不是识别错误）。",
        f"3. 黄色底纹 = 金额为 {_format_hits(hits)} 元的交易。",
        "4. 同一年份被拆成多段截图时，段与段的重叠部分按「日期+时间+金额+余额」去重；"
        "各分表按原图保留，汇总与全部明细已去重。",
        "5. 图片文件按内容 MD5 去重，内容完全相同的副本只算一张，见下方文件对照。",
    ]
    for note in notes:
        ws.cell(row=row, column=1, value=note).font = Font(size=9, bold=note == "核对说明")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="文件对照（重复文件 → 唯一截图）").font = Font(
        bold=True, size=10, color="1F4E79"
    )
    row += 1
    _headers(ws, row, ["工作表", "账单区间", "代表文件", "重复副本"])
    for page in all_pages:
        row += 1
        group = next((g for g in groups if page in g.pages), None)
        values = [
            page.name,
            group.period_label if group else "",
            page.source,
            "、".join(page.copies) or "—",
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = _CENTER if column == 1 else _LEFT

    for index, width in enumerate([12, 26, 34, 14, 15, 15, 14, 18, 13, 14], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"

    # ---------- 全部明细 ----------
    ws_all = workbook.create_sheet("全部明细(去重)")
    ws_all["A1"] = (
        f"全部交易明细（去重后 {len(records)} 笔；支出 {total_expense:,.2f} 元，收入 {total_income:,.2f} 元）"
    )
    ws_all["A1"].font = Font(bold=True, size=11, color="1F4E79")
    end = _write_records(ws_all, records, 2, with_source=True)
    _highlight(ws_all, records, 2, hits, with_source=True)
    row = end + 2
    ws_all.cell(row=row, column=2, value="支出合计").font = Font(bold=True)
    cell = ws_all.cell(row=row, column=4, value=total_expense)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    ws_all.cell(row=row, column=5, value="收入合计").font = Font(bold=True)
    cell = ws_all.cell(row=row, column=6, value=total_income)
    cell.number_format, cell.font = _MONEY, Font(bold=True)
    for index, width in enumerate(WIDTHS, 1):
        ws_all.column_dimensions[get_column_letter(index)].width = width
    ws_all.freeze_panes = "A3"

    # ---------- 每张唯一截图一个表 ----------
    used_names = {ws.title, ws_all.title}
    for page in all_pages:
        sheet_name = _unique_sheet_name(page.name, used_names)
        used_names.add(sheet_name)
        ws_page = workbook.create_sheet(sheet_name)
        page_records = sorted(page.records, key=lambda r: (r.date, r.time))
        expense, income = _totals(page_records)
        group = next((g for g in groups if page in g.pages), None)
        ws_page["A1"] = (
            f"{page.name}　来源：{page.source}"
            + (f"（副本 {'、'.join(page.copies)}）" if page.copies else "")
            + (f"　区间：{group.period_label}" if group else "")
        )
        ws_page["A1"].font = Font(bold=True, size=11, color="1F4E79")
        ws_page["A2"] = (
            f"共 {len(page_records)} 笔　支出 {expense:,.2f} 元　收入 {income:,.2f} 元"
            f"　黄色行为 {_format_hits(hits)} 元"
        )
        ws_page["A2"].font = Font(size=9, color="7F3F00", bold=True)
        _write_records(ws_page, page_records, 3, with_source=False)
        _highlight(ws_page, page_records, 3, hits, with_source=False)
        for index, width in enumerate(WIDTHS[:-1], 1):
            ws_page.column_dimensions[get_column_letter(index)].width = width
        ws_page.freeze_panes = "A4"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(out_path)
    except PermissionError as exc:  # 文件被 Excel 打开时 openpyxl 抛的就是这个
        raise PermissionError(f"无法写入 {out_path.name}，请先关闭同名文件") from exc
    return out_path


def _format_hits(hits: Iterable[float]) -> str:
    values = sorted({float(v) for v in hits})
    return " / ".join(f"{v:g}" for v in values) if values else "（无）"
